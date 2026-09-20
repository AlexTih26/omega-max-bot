"""Экспорт тестового реестра РУМЕКС в Excel."""

from __future__ import annotations

import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

_TIME_ZONE = ZoneInfo("Asia/Irkutsk")

_STATUS_LABELS = {
    "awaiting_accountant_review": "Ожидает проверки бухгалтера",
    "requires_correction": "Возвращена на исправление",
    "awaiting_er_sent": "Ожидает отметки в Контуре",
    "documents_ready": "Документы открыты диспетчеру",
    "documents_handed_to_driver": "Документы переданы водителю",
}


def _date(value: object) -> str:
    try:
        return datetime.fromtimestamp(float(value), tz=_TIME_ZONE).strftime("%d.%m.%Y")
    except (TypeError, ValueError, OSError):
        return ""


def _matches(shipment: dict, *, date_from: str, date_to: str, search: str) -> bool:
    loaded_date = datetime.fromtimestamp(float(shipment["loaded_at"]), tz=_TIME_ZONE).strftime("%Y-%m-%d")
    if date_from and loaded_date < date_from:
        return False
    if date_to and loaded_date > date_to:
        return False
    if not search:
        return True
    vehicle = (shipment.get("document_snapshot") or {}).get("vehicle") or {}
    values = [shipment.get("ttn_number"), shipment.get("registry_number"), vehicle.get("plate_tail"), vehicle.get("full_plate")]
    return search.lower() in " ".join(str(value or "") for value in values).lower()


def build_test_registry_workbook(
    shipments: list[dict], *, date_from: str = "", date_to: str = "", search: str = ""
) -> bytes:
    """Сформировать один Excel-лист с отфильтрованным реестром и всеми блоками."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Реестр РУМЕКС"
    headers = [
        "Дата погрузки",
        "№ ТТН",
        "ФИО водителя",
        "Гос. номер машины",
        "Буква",
        "Номер буквы",
        "Дата выдачи документа",
        "Последний статус",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    for shipment in shipments:
        if not _matches(shipment, date_from=date_from, date_to=date_to, search=search):
            continue
        snapshot = shipment.get("document_snapshot") or {}
        vehicle = snapshot.get("vehicle") or {}
        driver = snapshot.get("driver") or {}
        items = shipment.get("items") or [{}]
        for item in items:
            sheet.append([
                _date(shipment.get("loaded_at")),
                shipment.get("ttn_number") or shipment.get("registry_number") or "",
                driver.get("full_name") or "",
                vehicle.get("full_plate") or vehicle.get("plate_tail") or "",
                item.get("block_type_code") or item.get("letter") or "",
                item.get("block_number") or item.get("number") or "",
                _date(shipment.get("ttn_assigned_at")),
                _STATUS_LABELS.get(shipment.get("status"), shipment.get("status") or ""),
            ])

    sheet.freeze_panes = "A2"
    for index, width in enumerate((16, 22, 30, 22, 11, 16, 22, 34), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
