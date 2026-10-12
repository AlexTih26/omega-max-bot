"""Печать четырёх экземпляров утверждённой ТТН внутри новой Таксимо.

Источник формы неизменяем: для каждого экземпляра открывается отдельная копия
утверждённого бланка, в которую подставляются только данные рейса. Старый
контур РУМЕКС не читается и не изменяется: новая программа сама печатает ТТН.
"""

from __future__ import annotations

from copy import copy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import os
from pathlib import Path
import re
import tempfile
from typing import Any, Mapping

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from zoneinfo import ZoneInfo


APPROVED_TTN_TEMPLATE_PATH = (
    Path(__file__).resolve().parent / "rumex_templates" / "ТТН РУМЕКС — утверждённый шаблон.xlsx"
)
TTN_ARCHIVE_DIR = Path(__file__).resolve().parent.parent / "docs" / "taksimo-new" / "ttn-archive"
COPY_NUMBERS = (1, 2, 3, 4)
MIN_BLOCKS = 3
MAX_BLOCKS = 6

OMEGA_DETAILS = (
    "ООО «Омега-М», ИНН 5406829253, КПП 502401001, "
    "143405, Московская область, г. Красногорск, ул. Почтовая, д. 3"
)
RUMEX_DETAILS = (
    "ООО «РУМЕКС», ИНН 9728126848, КПП 201001001, "
    "364030, Чеченская Республика, г.о. город Грозный, "
    "г. Грозный, р-н Байсангуровский, ул. Сайханова, двлд. 222"
)
CARRIER_DETAILS = (
    "ООО «Комсомольская ТК», ИНН 2721252270, "
    "664025, Иркутская область, г. Иркутск, ул. Сурикова, д. 6, офис 2"
)
CARRIER_DETAILS_SECOND_PAGE = (
    "ООО «Комсомольская ТК», ИНН 2721252270\n"
    "664025, Иркутская область, г. Иркутск, ул. Сурикова, д. 6, офис 2"
)
OMEGA_DETAILS_SECOND_PAGE = (
    "ООО «Омега-М»\n"
    "ИНН 5406829253, КПП 502401001\n"
    "143405, Московская область, г. Красногорск, ул. Почтовая, д. 3"
)
PICKUP_LOCATION = (
    "Завод по производству тоннельной обделки на восточном\n"
    "портале Северомуйского тоннеля пгт. Северомуйск,\n"
    "расположенный Республика Бурятия, Муйский р-н,\n"
    "ГП «Северомуйское», пгт. Северомуйск"
)
DELIVERY_LOCATION = "ж/д станция Таксимо (код станции ЕСР 90440)"
PRODUCT_NAME = "БЛОК 9,7/8,8"

TTN_NUMBER_RE = re.compile(r"^ТТН №РМ-\d{4}-\d{6}$")


def _timezone() -> ZoneInfo:
    try:
        return ZoneInfo(os.getenv("TAKSIMO_NEW_TIMEZONE", "Asia/Irkutsk"))
    except Exception:
        return ZoneInfo("Asia/Irkutsk")


def _safe_filename_part(value: str, *, fallback: str) -> str:
    text = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', " ", (value or "").strip())
    text = re.sub(r"\s+", " ", text).strip(" .")
    return (text or fallback)[:160]


def _required_text(value: Any, label: str) -> str:
    text = str(value or "").strip()
    if not text:
        raise ValueError(f"В ТТН отсутствует: {label}")
    return text


def _copy_number(value: Any) -> int:
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError("Укажите экземпляр ТТН от 1 до 4") from None
    if number not in COPY_NUMBERS:
        raise ValueError("Укажите экземпляр ТТН от 1 до 4")
    return number


def _approved_number(shipment: Mapping[str, Any]) -> str:
    number = str(shipment.get("ttn_number") or "").strip()
    if not TTN_NUMBER_RE.fullmatch(number):
        raise ValueError("Для рейса ещё не выдан утверждённый номер ТТН")
    return number


def _loaded_datetime(shipment: Mapping[str, Any]) -> datetime:
    value = shipment.get("loaded_at")
    if isinstance(value, datetime):
        loaded = value
    else:
        try:
            timestamp = float(value)
        except (TypeError, ValueError):
            raise ValueError("У рейса не указана дата погрузки") from None
        if timestamp <= 0:
            raise ValueError("У рейса не указана дата погрузки")
        loaded = datetime.fromtimestamp(timestamp, _timezone())
    if loaded.tzinfo is None:
        loaded = loaded.replace(tzinfo=_timezone())
    return loaded.astimezone(_timezone())


def _surname_with_initials(full_name: str) -> str:
    text = " ".join(str(full_name or "").split())
    parts = text.split(" ")
    if len(parts) < 2:
        return text
    surname, rest = parts[0], " ".join(parts[1:])
    existing_initials = re.findall(r"[A-Za-zА-Яа-яЁё]\.", rest)
    if existing_initials and re.fullmatch(r"[A-Za-zА-Яа-яЁё.\s]+", rest):
        return surname + " " + "".join(existing_initials)
    initials = "".join(part[0].upper() + "." for part in parts[1:] if part)
    return surname + " " + initials if initials else surname


def _make_yellow_cells_white(sheet) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            fill = cell.fill
            if fill.fill_type != "solid":
                continue
            color = fill.fgColor
            rgb = str(color.rgb or "").upper() if color.type == "rgb" else ""
            if rgb.endswith("FFFF00"):
                cell.fill = PatternFill(fill_type="solid", fgColor="FFFFFFFF", bgColor="FFFFFFFF")


def _items(shipment: Mapping[str, Any]) -> list[dict]:
    items = [item for item in (shipment.get("items") or []) if isinstance(item, Mapping)]
    if not MIN_BLOCKS <= len(items) <= MAX_BLOCKS:
        raise ValueError(f"Утверждённая ТТН формируется только для рейса из {MIN_BLOCKS}-{MAX_BLOCKS} блоков")
    return items


def _total_weight(items: list[Mapping[str, Any]]) -> int:
    total = 0
    for item in items:
        try:
            total += int(item["weight_kg"])
        except (KeyError, TypeError, ValueError):
            raise ValueError("В ТТН отсутствует утверждённый вес одного из блоков") from None
    if total <= 0:
        raise ValueError("В ТТН отсутствует утверждённый вес блоков")
    return total


def _clear_cells(sheet, *references: str) -> None:
    for reference in references:
        sheet[reference] = None


def ttn_filename(shipment: Mapping[str, Any], *, copy_number: int) -> str:
    number = _approved_number(shipment)
    copy_value = _copy_number(copy_number)
    return _safe_filename_part(number, fallback="ТТН") + f" — Экземпляр № {copy_value}.xlsx"


def build_ttn_workbook(shipment: Mapping[str, Any], *, copy_number: int) -> bytes:
    """Сформировать один экземпляр утверждённого бланка ТТН по данным рейса."""
    if not APPROVED_TTN_TEMPLATE_PATH.is_file():
        raise ValueError("Утверждённый шаблон ТТН не загружен")
    number = _approved_number(shipment)
    copy_value = _copy_number(copy_number)
    loaded_at = _loaded_datetime(shipment)
    items = _items(shipment)
    total_weight_kg = _total_weight(items)

    snapshot = shipment.get("document_snapshot") or {}
    vehicle = snapshot.get("vehicle") or {}
    driver = snapshot.get("driver") or {}
    driver_name = _surname_with_initials(_required_text(driver.get("full_name"), "ФИО водителя"))
    driver_license = _required_text(driver.get("license_number"), "водительское удостоверение")
    vehicle_model = _required_text(vehicle.get("model"), "модель автомобиля")
    vehicle_plate = _required_text(vehicle.get("full_plate"), "государственный номер автомобиля")

    workbook = load_workbook(APPROVED_TTN_TEMPLATE_PATH, data_only=False)
    if "ТТН" not in workbook.sheetnames:
        raise ValueError("Утверждённый шаблон ТТН имеет неверную структуру")
    sheet = workbook["ТТН"]
    _make_yellow_cells_white(sheet)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 2
    sheet.page_setup.scale = None

    sheet["G5"] = loaded_at.date()
    sheet["AC5"] = number
    sheet["W6"] = str(copy_value)
    sheet["W8"] = "✓"
    sheet["A9"] = OMEGA_DETAILS
    sheet["BO9"] = RUMEX_DETAILS
    _clear_cells(sheet, "A11", "BO11")
    sheet["A14"] = OMEGA_DETAILS
    sheet["A16"] = DELIVERY_LOCATION

    for row_number, item in zip(range(20, 27), items, strict=False):
        code = _required_text(item.get("block_type_code"), "тип блока")
        block_number = _required_text(item.get("block_number"), "номер блока")
        sheet[f"A{row_number}"] = PRODUCT_NAME
        sheet[f"BF{row_number}"] = f"{code} {block_number}"
        sheet.row_dimensions[row_number].hidden = False
    for row_number in range(20 + len(items), 27):
        _clear_cells(sheet, f"A{row_number}", f"BF{row_number}")
        sheet.row_dimensions[row_number].hidden = True
    sheet["A28"] = f"{total_weight_kg:,} кг".replace(",", " ")
    _clear_cells(sheet, "A30", "BE30", "A33", "A35")
    sheet["B38"] = "-"

    sheet["B43"] = CARRIER_DETAILS
    sheet["BE43"] = driver_name
    sheet["CG43"] = driver_license
    _clear_cells(sheet, "CX43")
    sheet["B46"] = vehicle_model
    sheet["BE46"] = vehicle_plate
    sheet["A49"] = "Тип владения: собственность"

    sheet["A52"] = PICKUP_LOCATION
    sheet["A54"] = RUMEX_DETAILS
    sheet["A56"] = PICKUP_LOCATION
    sheet["A58"] = loaded_at.date()
    sheet["AC58"] = loaded_at.time().replace(second=0, microsecond=0)
    sheet["A60"] = f"{total_weight_kg:,}".replace(",", " ") + " кг, расчётная масса"
    sheet["A62"] = len(items)
    sheet["BE62"] = "без тары"
    sheet["A64"] = "-"
    _clear_cells(sheet, "BE58", "CH58")

    sheet["A74"] = DELIVERY_LOCATION
    _clear_cells(sheet, "A76", "Z76", "BE76", "A78", "BE78", "AF80", "BE80", "B90", "B96", "B97")

    for reference, value in (
        ("B92", CARRIER_DETAILS_SECOND_PAGE),
        ("BF92", OMEGA_DETAILS_SECOND_PAGE),
        ("BF96", OMEGA_DETAILS_SECOND_PAGE),
    ):
        cell = sheet[reference]
        cell.value = value
        font = copy(cell.font)
        font.sz = 8
        cell.font = font
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        alignment.vertical = "center"
        cell.alignment = alignment

    for row_number, height in ((9, 30), (54, 18), (92, 48), (96, 48)):
        sheet.row_dimensions[row_number].height = height

    calculation = workbook.calculation
    calculation.fullCalcOnLoad = True
    calculation.forceFullCalc = True
    calculation.calcMode = "auto"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def archive_ttn_copies(shipment: Mapping[str, Any]) -> list[dict[str, Any]]:
    """Сохранить неизменяемые экземпляры ТТН рейса в отдельный архив."""
    trip_public_id = str(shipment.get("public_id") or "").strip()
    if not trip_public_id:
        raise ValueError("Не указан рейс для архива ТТН")
    TTN_ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
    archived: list[dict[str, Any]] = []
    for copy_number in COPY_NUMBERS:
        content = build_ttn_workbook(shipment, copy_number=copy_number)
        digest = sha256(content).hexdigest()
        filename = f"trip-{trip_public_id}-ttn-copy-{copy_number}.xlsx"
        destination = TTN_ARCHIVE_DIR / filename
        if destination.exists():
            if sha256(destination.read_bytes()).hexdigest() != digest:
                raise ValueError("Архивный экземпляр ТТН не совпадает с выданным документом")
        else:
            with tempfile.NamedTemporaryFile(dir=TTN_ARCHIVE_DIR, delete=False) as temporary:
                temporary.write(content)
                temp_path = Path(temporary.name)
            try:
                temp_path.replace(destination)
            finally:
                temp_path.unlink(missing_ok=True)
        archived.append({
            "copy_number": copy_number,
            "filename": filename,
            "sha256": digest,
            "size_bytes": len(content),
        })
    return archived


def build_copies_for_store(shipment: Mapping[str, Any], copy_number: int) -> bytes:
    """Точка входа для store: собрать экземпляр и проверить, что он строится."""
    return build_ttn_workbook(shipment, copy_number=copy_number)
