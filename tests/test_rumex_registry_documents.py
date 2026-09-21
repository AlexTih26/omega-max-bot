"""Tests for approved XLSX generation in the isolated RUMEX test flow."""

from __future__ import annotations

import hashlib
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
BOT = ROOT / "fotonych-bot"
if str(BOT) not in sys.path:
    sys.path.insert(0, str(BOT))

import rumex_registry_documents as documents  # noqa: E402


class RumexRegistryDocumentsTests(unittest.TestCase):
    def test_approved_ttn_preserves_template_and_creates_all_four_copies(self) -> None:
        shipment = {
            "registry_number": "РМ-2026-000123",
            "ttn_number": "ТТН №РМ-2026-000001",
            "loaded_at": 1_767_225_600.0,
            "items": [
                {"block_type_code": "K", "block_number": "7741", "weight_kg": 3850},
                {"block_type_code": "A", "block_number": "3611", "weight_kg": 7780},
                {"block_type_code": "A", "block_number": "3612", "weight_kg": 7780},
            ],
            "document_snapshot": {
                "vehicle": {"plate_tail": "553", "full_plate": "К553НХ 138", "model": "FAW J6"},
                "driver": {
                    "full_name": "Иванов Иван Иванович",
                    "license_number": "38 12 123456",
                },
            },
        }
        source_hash = hashlib.sha256(documents.APPROVED_TTN_TEMPLATE_PATH.read_bytes()).hexdigest()
        self.assertEqual(
            source_hash,
            "02d8a20a0d9b420b69e7fce5581038e9c80526fa83d5501b98a63dc46270f093",
        )

        for copy_number in (1, 2, 3, 4):
            content = documents.build_test_ttn_workbook(shipment, copy_number=copy_number)
            self.assertGreater(len(content), 10_000)
            with tempfile.NamedTemporaryFile(suffix=".xlsx") as generated:
                generated.write(content)
                generated.flush()
                workbook = load_workbook(generated.name, data_only=False)
            sheet = workbook["ТТН"]
            self.assertEqual(sheet["AC5"].value, "ТТН №РМ-2026-000001")
            self.assertEqual(sheet["W6"].value, str(copy_number))
            self.assertEqual(sheet["BF20"].value, "K 7741")
            self.assertEqual(sheet["BF21"].value, "A 3611")
            self.assertEqual(sheet["BF22"].value, "A 3612")
            self.assertEqual([sheet[f"BF{row}"].value for row in range(23, 27)], [None, None, None, None])
            self.assertEqual([sheet.row_dimensions[row].hidden for row in range(20, 23)], [False, False, False])
            self.assertEqual([sheet.row_dimensions[row].hidden for row in range(23, 27)], [True, True, True, True])
            self.assertEqual(sheet["A28"].value, "19 410 кг")
            self.assertEqual(sheet["A60"].value, "19 410 кг, расчётная масса")
            self.assertEqual(sheet["A62"].value, 3)
            self.assertEqual(sheet["BE62"].value, "без тары")
            self.assertEqual(sheet["CG43"].value, "38 12 123456")
            self.assertEqual(sheet["BE43"].value, "Иванов И.И.")
            self.assertEqual(
                sheet["A9"].value,
                "ООО «Омега-М», ИНН 5406829253, КПП 502401001, "
                "143405, Московская область, г. Красногорск, ул. Почтовая, д. 3",
            )
            self.assertEqual(
                sheet["B92"].value,
                "ООО «Комсомольская ТК», ИНН 2721252270\n"
                "664025, Иркутская область, г. Иркутск, ул. Сурикова, д. 6, офис 2",
            )
            self.assertEqual(
                sheet["BF92"].value,
                "ООО «Омега-М»\n"
                "ИНН 5406829253, КПП 502401001\n"
                "143405, Московская область, г. Красногорск, ул. Почтовая, д. 3",
            )
            self.assertEqual(sheet["BF92"].font.sz, 8)
            self.assertTrue(sheet["BF92"].alignment.wrap_text)
            self.assertEqual(sheet.row_dimensions[92].height, 48)
            self.assertIn("ООО «РУМЕКС», ИНН", sheet["BO9"].value)
            self.assertNotIn("ОГРН", sheet["A9"].value)
            self.assertNotIn("ОГРН", sheet["BO9"].value)
            self.assertNotIn("\n", sheet["B43"].value)
            self.assertIsNone(sheet["BE58"].value)
            self.assertIsNone(sheet["CH58"].value)
            self.assertIsNone(sheet["A78"].value)
            self.assertIsNone(sheet["BE78"].value)
            self.assertIsNone(sheet["AF80"].value)
            self.assertIsNone(sheet["BE80"].value)
            self.assertEqual(sheet["BE56"].value, "=G5")
            self.assertEqual(sheet["CO66"].value, "=BE43")
            self.assertEqual(sheet["CO82"].value, "=CO66")
            self.assertEqual(str(sheet.print_area), "'ТТН'!$A$1:$DG$101")
            self.assertEqual(sheet.page_setup.orientation, "portrait")
            self.assertEqual(sheet.page_setup.paperSize, 9)
            self.assertEqual(sheet.page_setup.fitToWidth, 1)
            self.assertEqual(sheet.page_setup.fitToHeight, 2)
            self.assertIsNone(sheet.page_setup.scale)
            for row in sheet.iter_rows():
                for cell in row:
                    fill = cell.fill
                    rgb = str(fill.fgColor.rgb or "").upper() if fill.fgColor.type == "rgb" else ""
                    self.assertFalse(fill.fill_type == "solid" and rgb.endswith("FFFF00"), cell.coordinate)
            self.assertEqual(
                documents.test_ttn_filename(shipment, copy_number=copy_number),
                f"ТТН №РМ-2026-000001 — Экземпляр № {copy_number}.xlsx",
            )

        self.assertEqual(
            hashlib.sha256(documents.APPROVED_TTN_TEMPLATE_PATH.read_bytes()).hexdigest(), source_hash
        )

    def test_approved_ttn_requires_assigned_number_and_valid_copy(self) -> None:
        shipment = {"ttn_number": "", "loaded_at": 1, "items": []}
        with self.assertRaisesRegex(ValueError, "не выдан утверждённый номер"):
            documents.build_test_ttn_workbook(shipment, copy_number=1)
        with self.assertRaisesRegex(ValueError, "экземпляр ТТН от 1 до 4"):
            documents.test_ttn_filename({"ttn_number": "ТТН №РМ-2026-000001"}, copy_number=5)

    def test_surname_with_initials_keeps_compact_name_and_shortens_full_name(self) -> None:
        self.assertEqual(documents._surname_with_initials("Иванов Иван Иванович"), "Иванов И.И.")
        self.assertEqual(documents._surname_with_initials("Бадрянов А.С."), "Бадрянов А.С.")
        self.assertEqual(documents._surname_with_initials("Петров Пётр"), "Петров П.")


if __name__ == "__main__":
    unittest.main()
