"""Проверки рейса, ТТН и экспедиторской расписки без подключения к MySQL."""

from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
import sys
import types
import unittest
from unittest.mock import patch


ROOT = Path(__file__).resolve().parent.parent
BOT = ROOT / "fotonych-bot"
if str(BOT) not in sys.path:
    sys.path.insert(0, str(BOT))

if "dotenv" not in sys.modules:
    dotenv = types.ModuleType("dotenv")
    dotenv.load_dotenv = lambda *_args, **_kwargs: False
    sys.modules["dotenv"] = dotenv

import taksimo_new_store as store  # noqa: E402
import taksimo_new_receipts as receipts  # noqa: E402
import taksimo_new_ttn as ttn  # noqa: E402


DISPATCHER = {"id": 11, "role": "dispatcher", "name": "Диспетчер"}
ACCOUNTANT = {"id": 12, "role": "accountant1", "name": "Бухгалтер 1"}
OPERATOR = {"id": 7, "role": "operator1", "name": "Оператор 1"}
ADMIN = {"id": 10, "role": "admin", "name": "Администратор"}


class _Cursor:
    def __init__(self) -> None:
        self.executed: list[tuple[str, object]] = []
        self.lastrowid = 1
        self._row: dict | None = None
        self._rows: list[dict] = []

    def __enter__(self) -> "_Cursor":
        return self

    def __exit__(self, _type, _value, _traceback) -> None:
        return None

    def execute(self, statement: str, parameters: object = None) -> None:
        query = " ".join(statement.split())
        self.executed.append((query, parameters))
        self._row = None
        self._rows = []
        if query.startswith("SELECT letter, weight_kg"):
            self._row = {"letter": "A", "weight_kg": 7780, "price_rub": 1200}
        elif "COALESCE(MAX(sequence_number)" in query:
            self._row = {"next_sequence": 1, "trip_year": 2026}

    def fetchone(self) -> dict | None:
        return self._row

    def fetchall(self) -> list[dict]:
        return self._rows


class _Connection:
    def __init__(self, cursor: _Cursor) -> None:
        self._cursor = cursor

    def cursor(self) -> _Cursor:
        return self._cursor


class TripRulesTests(unittest.TestCase):
    def _run(self, function, operator, **kwargs):
        cursor = _Cursor()
        connection = _Connection(cursor)

        @contextmanager
        def transaction():
            yield connection

        with patch.object(store, "transaction", transaction), patch.object(
            store, "utc_now", return_value=datetime(2026, 1, 1, 12, 0, 0)
        ):
            return function(operator=operator, **kwargs), cursor

    def test_operator_cannot_create_trip(self) -> None:
        cursor = _Cursor()
        connection = _Connection(cursor)

        @contextmanager
        def transaction():
            yield connection

        with patch.object(store, "transaction", transaction):
            with self.assertRaises(PermissionError):
                store.create_trip({"vehicle_plate": "А123ВС75", "blocks": []}, operator=OPERATOR)

    def test_accountant_cannot_create_trip(self) -> None:
        cursor = _Cursor()
        connection = _Connection(cursor)

        @contextmanager
        def transaction():
            yield connection

        with patch.object(store, "transaction", transaction):
            with self.assertRaises(PermissionError):
                store.create_trip({"vehicle_plate": "А123ВС75", "blocks": []}, operator=ACCOUNTANT)

    def test_dispatcher_roles_include_admin(self) -> None:
        self.assertIn("dispatcher", store.DISPATCHER_ROLES)
        self.assertIn("admin", store.DISPATCHER_ROLES)

    def test_trip_rejects_unknown_block_letter(self) -> None:
        cursor = _Cursor()
        connection = _Connection(cursor)

        @contextmanager
        def transaction():
            yield connection

        with patch.object(store, "transaction", transaction):
            with self.assertRaises(ValueError):
                store.create_trip(
                    {"vehicle_plate": "А123ВС75", "blocks": [{"letter": "Z", "number": "1"}]},
                    operator=DISPATCHER,
                )


class TtnPrintTests(unittest.TestCase):
    def _shipment(self, **overrides):
        base = {
            "public_id": "3f2a6c1e-0000-0000-0000-000000000001",
            "ttn_number": "ТТН №РМ-2026-000007",
            "loaded_at": 1760000000.0,
            "items": [
                {"block_type_code": "A", "block_number": "101", "weight_kg": 7780, "price_rub": 1200},
                {"block_type_code": "B", "block_number": "102", "weight_kg": 7900, "price_rub": 1300},
                {"block_type_code": "C", "block_number": "103", "weight_kg": 8080, "price_rub": 1400},
            ],
            "document_snapshot": {
                "vehicle": {"model": "КАМАЗ 65115", "full_plate": "А123ВС75"},
                "driver": {"full_name": "Иванов Иван Иванович", "license_number": "75 АА 123456"},
            },
        }
        base.update(overrides)
        return base

    def test_document_builds_for_each_of_four_copies(self) -> None:
        for copy_number in (1, 2, 3, 4):
            content = ttn.build_ttn_workbook(self._shipment(), copy_number=copy_number)
            self.assertGreater(len(content), 10_000, f"экземпляр {copy_number} пуст")

    def test_filename_contains_number_and_copy(self) -> None:
        name = ttn.ttn_filename(self._shipment(), copy_number=3)
        self.assertIn("ТТН №РМ-2026-000007", name)
        self.assertIn("Экземпляр № 3", name)

    def test_ttn_requires_issued_number(self) -> None:
        with self.assertRaises(ValueError):
            ttn.build_ttn_workbook(self._shipment(ttn_number=""), copy_number=1)

    def test_ttn_rejects_fewer_than_three_blocks(self) -> None:
        shipment = self._shipment(items=self._shipment()["items"][:2])
        with self.assertRaises(ValueError):
            ttn.build_ttn_workbook(shipment, copy_number=1)

    def test_ttn_rejects_more_than_six_blocks(self) -> None:
        items = self._shipment()["items"] * 3
        with self.assertRaises(ValueError):
            ttn.build_ttn_workbook(self._shipment(items=items), copy_number=1)

    def test_ttn_requires_driver_license(self) -> None:
        snapshot = {
            "vehicle": {"model": "КАМАЗ", "full_plate": "А123ВС75"},
            "driver": {"full_name": "Иванов Иван", "license_number": ""},
        }
        with self.assertRaises(ValueError):
            ttn.build_ttn_workbook(self._shipment(document_snapshot=snapshot), copy_number=1)

    def test_ttn_rejects_copy_outside_one_to_four(self) -> None:
        for bad in (0, 5, "нет"):
            with self.assertRaises(ValueError):
                ttn.build_ttn_workbook(self._shipment(), copy_number=bad)

    def test_driver_name_is_printed_with_initials(self) -> None:
        from io import BytesIO
        from openpyxl import load_workbook

        content = ttn.build_ttn_workbook(self._shipment(), copy_number=1)
        sheet = load_workbook(BytesIO(content))["ТТН"]
        self.assertEqual(sheet["BE43"].value, "Иванов И.И.")
        self.assertEqual(sheet["BE46"].value, "А123ВС75")
        self.assertEqual(sheet["AC5"].value, "ТТН №РМ-2026-000007")
        self.assertEqual(sheet["A62"].value, 3)

    def test_total_weight_is_sum_of_blocks(self) -> None:
        from io import BytesIO
        from openpyxl import load_workbook

        content = ttn.build_ttn_workbook(self._shipment(), copy_number=2)
        sheet = load_workbook(BytesIO(content))["ТТН"]
        self.assertEqual(sheet["A28"].value, "23 760 кг")

    def test_unused_block_rows_are_hidden(self) -> None:
        from io import BytesIO
        from openpyxl import load_workbook

        content = ttn.build_ttn_workbook(self._shipment(), copy_number=1)
        sheet = load_workbook(BytesIO(content))["ТТН"]
        self.assertFalse(sheet.row_dimensions[20].hidden)
        self.assertTrue(sheet.row_dimensions[23].hidden)


class ReceiptRulesTests(unittest.TestCase):
    def test_operator_cannot_create_receipt(self) -> None:
        cursor = _Cursor()
        connection = _Connection(cursor)

        @contextmanager
        def transaction():
            yield connection

        with patch.object(receipts, "transaction", transaction):
            with self.assertRaises(PermissionError):
                receipts.create_receipt("3f2a6c1e-0000-0000-0000-000000000001", operator=OPERATOR)

    def test_dispatcher_cannot_create_receipt(self) -> None:
        cursor = _Cursor()
        connection = _Connection(cursor)

        @contextmanager
        def transaction():
            yield connection

        with patch.object(receipts, "transaction", transaction):
            with self.assertRaises(PermissionError):
                receipts.create_receipt("3f2a6c1e-0000-0000-0000-000000000001", operator=DISPATCHER)

    def test_receipt_blocks_text_needs_accountant_role(self) -> None:
        cursor = _Cursor()
        connection = _Connection(cursor)

        @contextmanager
        def transaction():
            yield connection

        with patch.object(receipts, "transaction", transaction):
            with self.assertRaises(PermissionError):
                receipts.receipt_blocks_text("3f2a6c1e-0000-0000-0000-000000000001", operator=OPERATOR)

    def test_conflict_error_is_a_value_error(self) -> None:
        self.assertTrue(issubclass(receipts.TaksimoNewConflictError, ValueError))


class TtnCopyRulesTests(unittest.TestCase):
    def test_builder_receives_copy_number_and_returns_bytes(self) -> None:
        seen: list[int] = []

        def builder(_shipment, copy_number):
            seen.append(copy_number)
            return b"x" * 32

        self.assertEqual(sorted(ttn.COPY_NUMBERS), [1, 2, 3, 4])
        self.assertEqual(len(builder({}, 1)), 32)
        self.assertEqual(seen, [1])


if __name__ == "__main__":
    unittest.main()
